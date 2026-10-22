# -*- coding: utf-8 -*-
import openpyxl, json, sys
sys.stdout.reconfigure(encoding="utf-8")
from _paths import SIE, SYS, PLAN
# Lee el PLANIFICADO real (portable). Si Excel lo tiene abierto/bloqueado,
# copialo antes a  _sistema/PLAN_copy.xlsx  y apuntá p a esa copia.
import os
p = SYS + "/PLAN_copy.xlsx" if os.path.exists(SYS + "/PLAN_copy.xlsx") else PLAN
wb = openpyxl.load_workbook(p, data_only=True); ws = wb["PLANIFICADO"]
CAPS = {"A":"A · Trabajos preliminares","C":"C · Estructura H°A°","D":"D · Estructura metálica + cubierta",
        "F":"F · Canaletas y cerramientos","G":"G · Mampostería, pisos y terminaciones",
        "H":"H · Cámara de recolección residuos","C.I":"CI · Costes indirectos"}
def norm(c): return "CI" if c in ("C.I","CI") else c
def num(v):
    try: return round(float(v))
    except: return 0
def s(v): return str(v).strip() if v is not None else ""

def es_agregador(r):
    """Renglones 'IN.' sin datos en M/O (Cantidad Objetivo/Total Gs) son un
    sub-grupo agregado (ej. 'Fabricación de Estructura tipo Cabriada...',
    'Fabricación de VIGA RIOSTRA...') — tienen su propio objetivo en D/I
    (CanObj/ImpObj, igual que una fila de partida) y encabezan sus propios
    insumos de detalle debajo. Se promueven a PARTIDA propia (no quedan
    anidados como insumo de la partida madre) para poder cargarles compras
    e imputarles objetivo directamente."""
    return ws.cell(r,13).value is None and ws.cell(r,15).value is None

def insumo_leaf(r, cod, desc):
    return {"cod":cod,"desc":desc,"cant":ws.cell(r,13).value,"unidad":s(ws.cell(r,14).value),
            "total":num(ws.cell(r,15).value),"tipo":s(ws.cell(r,17).value)}

partidas=[]; cap=None; cur=None; ci=None; pid=0
for r in range(2, ws.max_row+1):
    cod=s(ws.cell(r,1).value); desc=s(ws.cell(r,2).value)
    if cod in CAPS: cap=norm(cod); cur=None; continue
    if not cod: continue
    if cap=="CI":
        if ci is None:
            pid+=1; ci={"id":f"P{pid:02d}","cod":"CI.GEN","cap":"CI","cap_nom":CAPS["C.I"],
                        "desc":"Costes indirectos","obj":53550000,"insumos":[]}
            partidas.append(ci); cur=ci
        cur["insumos"].append(insumo_leaf(r, cod, desc))
        continue
    if cod.startswith("RU.") or cod=="A.1":
        pid+=1; cur={"id":f"P{pid:02d}","cod":cod,"cap":cap,"cap_nom":CAPS[cap],
                     "desc":desc,"obj":num(ws.cell(r,9).value),"insumos":[]}
        partidas.append(cur); continue
    if cod.startswith("IN.") or cod.startswith("INSBOP."):
        if es_agregador(r):
            pid+=1; cur={"id":f"P{pid:02d}","cod":cod,"cap":cap,"cap_nom":CAPS[cap],
                         "desc":desc,"obj":num(ws.cell(r,9).value),"insumos":[]}
            partidas.append(cur); continue
        if cur is not None:
            cur["insumos"].append(insumo_leaf(r, cod, desc))

for capk in ["A","C","D","F","G","H"]:
    if not any(x["cod"]==f"{capk}.GEN" for x in partidas):
        pid+=1; partidas.append({"id":f"P{pid:02d}","cod":f"{capk}.GEN","cap":capk,"cap_nom":CAPS[capk],
                                 "desc":"(genérico del capítulo)","obj":0,"insumos":[]})

# PRESERVAR ADICIONALES: los capítulos AD* (ej. AD1 acometida eléctrica) NO están en el
# PLANIFICADO — viven solo en el JSON. Al re-extraer se arrastran del archivo existente.
try:
    _prev = json.load(open(SIE + "/tablero_economico/partidas_insumos.json", encoding="utf-8"))
    for _p in _prev["partidas"]:
        if _p["cap"].startswith("AD"):
            pid += 1; _p["id"] = f"P{pid:02d}"; partidas.append(_p)
            CAPS.setdefault(_p["cap"], _p["cap_nom"])
except FileNotFoundError:
    pass

ni=sum(len(x["insumos"]) for x in partidas)
print(f"Partidas: {len(partidas)} · Insumos totales: {ni}")
for x in partidas[:7]:
    print(f"  {x['id']} {x['cap']} {x['cod']} · {x['desc'][:38]}  (insumos: {len(x['insumos'])})")
    for ins in x["insumos"][:3]:
        print(f"       - {ins['cod']}  {str(ins['desc'])[:34]:36} cant {ins['cant']} {ins['unidad']}  ${ins['total']:,}")

full=SIE + "/tablero_economico/partidas_insumos.json"
_caps_out = {norm(k):v for k,v in CAPS.items()}
json.dump({"partidas":partidas,"caps":_caps_out},
          open(full,"w",encoding="utf-8"),ensure_ascii=False,indent=1)
app=[{"id":x["id"],"cod":x["cod"],"cap":x["cap"],"desc":x["desc"],
      "ins":[{"cod":i["cod"],"desc":i["desc"],"cant":i["cant"],"un":i["unidad"]} for i in x["insumos"]]} for x in partidas]
appout=SYS + "/partidas_app.json"
json.dump({"partidas":app,"caps":_caps_out},open(appout,"w",encoding="utf-8"),ensure_ascii=False)
print("guardado partidas_insumos.json (full) y partidas_app.json (sin plata)")
